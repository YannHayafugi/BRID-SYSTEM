"""Extração local de texto do Termo de Referência (PDF, DOCX ou TXT)."""
from pathlib import Path


def extrair_texto(caminho: str | Path) -> str:
    caminho = Path(caminho)
    ext = caminho.suffix.lower()
    if ext == ".pdf":
        return _extrair_pdf(caminho)
    if ext == ".docx":
        return _extrair_docx(caminho)
    if ext in (".txt", ".md"):
        return caminho.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Formato não suportado: {ext}. Use PDF, DOCX ou TXT.")


def _extrair_pdf(caminho: Path) -> str:
    return analisar_pdf(caminho)["texto"]


def analisar_pdf(caminho: str | Path) -> dict:
    """Extrai texto e estatísticas do PDF para decidir se OCR é necessário.

    Retorna: {"texto", "total_paginas", "paginas_com_texto"}
    """
    import pdfplumber

    partes, com_texto = [], 0
    with pdfplumber.open(caminho) as pdf:
        total = len(pdf.pages)
        for pagina in pdf.pages:
            texto = pagina.extract_text() or ""
            if len(texto.strip()) > 50:
                com_texto += 1
            if texto.strip():
                partes.append(texto)
    return {"texto": "\n\n".join(partes), "total_paginas": total, "paginas_com_texto": com_texto}


def precisa_ocr(info: dict) -> bool:
    """PDF digitalizado: quase nenhum texto extraível nas páginas."""
    total = info["total_paginas"]
    if total == 0:
        return False
    if len(info["texto"].strip()) < 100:
        return True
    return info["paginas_com_texto"] / total < 0.3


def _extrair_docx(caminho: Path) -> str:
    import docx

    doc = docx.Document(str(caminho))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabela in doc.tables:
        for linha in tabela.rows:
            celulas = [c.text.strip() for c in linha.cells if c.text.strip()]
            if celulas:
                partes.append(" | ".join(celulas))
    return "\n".join(partes)


def limpar_texto(texto: str, max_chars: int = 60000) -> str:
    """Remove linhas repetidas/vazias e trunca para economizar tokens."""
    linhas, anteriores = [], set()
    for linha in texto.splitlines():
        linha = linha.strip()
        if not linha:
            continue
        # remove cabeçalhos/rodapés repetidos (aparecem em toda página)
        if linha in anteriores and len(linha) < 80:
            continue
        anteriores.add(linha)
        linhas.append(linha)
    resultado = "\n".join(linhas)
    return resultado[:max_chars]
