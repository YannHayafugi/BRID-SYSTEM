"""Gerador de ofícios FIA (.docx) — porte do repo OFICIO-INICAL-SECURITIZACAO.

Módulo isolado: falha aqui não afeta a geração de propostas.
Geração 100% local (python-docx), sem chamadas de IA.
"""
import os
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Twips
from fastapi import APIRouter, Header, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field

router = APIRouter()

ASSETS_DIR = Path(__file__).resolve().parent.parent / "assets"


# ---------- modelos ----------

class ItemNorma(BaseModel):
    text: str
    level: int = 0  # 0 = numerado (1, 2...), 1 = subitem (a, b, c...)


class OficioData(BaseModel):
    destinatario_nome: str
    destinatario_endereco: list[str] = Field(default_factory=list)
    numero_contrato: str
    assunto: str
    ac_nome: str = ""
    ac_cargo: str = ""
    saudacao: str
    paragrafo_abertura: str
    etapas: str = ""
    paragrafo_etapas: str = ""
    introducao_lista_dados: str = ""
    itens_lista_dados: list[str] = Field(default_factory=list)
    paragrafo_transicao_normas: str = ""
    introducao_lista_normas: str = ""
    itens_lista_normas: list[ItemNorma] = Field(default_factory=list)
    paragrafo_pos_lista: str = ""
    paragrafo_reforcamos: str = ""
    paragrafo_tempestividade: str = ""
    paragrafo_teams: str = ""
    paragrafo_sharepoint_prazo: str = ""
    sharepoint_exclusividade1: str = ""
    sharepoint_exclusividade2: str = ""
    sharepoint_exclusividade3: str = ""
    paragrafo_reuniao: str = ""
    paragrafo_encerramento: str = ""
    cidade_emissao: str = "São Paulo"
    data_extenso: str = ""
    assinatura: str = "FUNDAÇÃO INSTITUTO DE ADMINISTRAÇÃO – FIA"


# ---------- helpers ----------

_ROMANOS = [(1000, "m"), (900, "cm"), (500, "d"), (400, "cd"), (100, "c"), (90, "xc"),
            (50, "l"), (40, "xl"), (10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]


def _romano(n: int) -> str:
    resultado = ""
    for valor, letra in _ROMANOS:
        while n >= valor:
            resultado += letra
            n -= valor
    return resultado


def _para(doc, spacing_after=12, line=1.15):
    p = doc.add_paragraph()
    f = p.paragraph_format
    f.space_after = Pt(spacing_after)
    f.line_spacing = line
    return p


def _body(doc, texto: str):
    p = _para(doc)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Twips(720)
    p.add_run(texto)
    return p


def _item_lista(doc, prefixo: str, texto: str, indent_esq: int):
    p = _para(doc, spacing_after=4)
    f = p.paragraph_format
    f.left_indent = Twips(indent_esq)
    f.first_line_indent = Twips(-360)
    f.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(f"{prefixo}\t{texto}")
    return p


# ---------- montagem do documento ----------

def montar_oficio(d: OficioData) -> bytes:
    doc = Document()

    # fonte padrão Garamond 11pt
    estilo = doc.styles["Normal"]
    estilo.font.name = "Garamond"
    estilo.font.size = Pt(11)

    sec = doc.sections[0]
    sec.top_margin = Twips(1800)
    sec.bottom_margin = Twips(1800)
    sec.left_margin = Twips(1440)
    sec.right_margin = Twips(1440)

    # cabeçalho: logo FIA centralizado no topo
    p_head = sec.header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.add_run().add_picture(str(ASSETS_DIR / "header_logo.jpg"),
                                 width=Inches(1.48), height=Inches(0.385))

    # rodapé: barra com dados de contato, largura total
    p_foot = sec.footer.paragraphs[0]
    p_foot.paragraph_format.left_indent = Twips(-1440)
    p_foot.add_run().add_picture(str(ASSETS_DIR / "footer_bar.jpg"), width=Inches(8.5), height=Inches(0.8))

    # destinatário
    p = _para(doc, spacing_after=10)
    p.add_run(d.destinatario_nome).bold = True
    for i, linha in enumerate(d.destinatario_endereco):
        _para(doc, spacing_after=15 if i == len(d.destinatario_endereco) - 1 else 2).add_run(linha)

    # Ref / Assunto
    p = _para(doc, spacing_after=2)
    r = p.add_run("Ref.: ")
    r.bold = True
    r.underline = True
    r = p.add_run(f"CONTRATO Nº {d.numero_contrato} – Fundação Instituto de Administração.")
    r.underline = True
    p = _para(doc, spacing_after=15)
    r = p.add_run("Assunto: ")
    r.bold = True
    r.underline = True
    p.add_run(d.assunto).underline = True

    # A/C
    if d.ac_nome:
        p = _para(doc, spacing_after=15)
        r = p.add_run(f"A/C de {d.ac_nome}")
        r.bold = True
        if d.ac_cargo:
            r = p.add_run(" – ")
            r.bold = True
            r = p.add_run(d.ac_cargo)
            r.bold = True
            r.italic = True
        p.add_run(";").bold = True

    # corpo
    _body(doc, f"{d.saudacao}, {d.paragrafo_abertura}")
    if d.paragrafo_etapas:
        _body(doc, d.paragrafo_etapas.replace("{ETAPAS}", d.etapas))
    if d.introducao_lista_dados:
        _body(doc, d.introducao_lista_dados)
    for n, item in enumerate(d.itens_lista_dados, start=1):
        _item_lista(doc, f"{_romano(n)}.", item, 720)

    if d.paragrafo_transicao_normas:
        _body(doc, d.paragrafo_transicao_normas)
    if d.introducao_lista_normas:
        _body(doc, d.introducao_lista_normas)
    num, letra = 0, 0
    for item in d.itens_lista_normas:
        if item.level == 0:
            num += 1
            letra = 0
            _item_lista(doc, f"{num}.", item.text, 720)
        else:
            letra += 1
            _item_lista(doc, f"{chr(96 + letra)})", item.text, 1440)

    for texto in (d.paragrafo_pos_lista, d.paragrafo_reforcamos, d.paragrafo_tempestividade,
                  d.paragrafo_teams, d.paragrafo_sharepoint_prazo,
                  d.sharepoint_exclusividade1, d.sharepoint_exclusividade2,
                  d.sharepoint_exclusividade3, d.paragrafo_reuniao, d.paragrafo_encerramento):
        if texto:
            _body(doc, texto)

    # fecho
    p = _para(doc, spacing_after=10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Atenciosamente,")
    p = _para(doc, spacing_after=30)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"{d.cidade_emissao}, {d.data_extenso}.")
    p = _para(doc)
    r = p.add_run(d.assinatura)
    r.bold = True
    r.italic = True

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------- endpoint ----------

def _senha_ok(senha: str) -> bool:
    esperada = os.getenv("APP_SENHA", "")
    return not esperada or senha == esperada


@router.post("/api/oficio")
def gerar_oficio(dados: OficioData, x_senha: str | None = Header(default=None)):
    if not _senha_ok(x_senha or ""):
        raise HTTPException(401, "Não autorizado. Faça login novamente.")
    try:
        conteudo = montar_oficio(dados)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(500, f"Falha ao gerar o ofício: {e}") from e
    nome = f"Oficio - Contrato {dados.numero_contrato}.docx".replace("/", "-")
    return StreamingResponse(
        BytesIO(conteudo),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{nome}"'},
    )
