"""Gera os templates DOCX iniciais (proposta.docx e resumo.docx).

Rode uma vez: python scripts/build_templates.py
Depois, edite os .docx no Word à vontade — mantenha as tags {{ ... }} / {% ... %}.
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt

TEMPLATES = Path(__file__).resolve().parent.parent / "templates"
TEMPLATES.mkdir(exist_ok=True)


def _base_doc():
    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(11)
    return doc


def build_proposta():
    doc = _base_doc()
    doc.add_heading("{{ titulo }}", 0)
    doc.add_paragraph("Cliente: {{ orgao_cliente }}")
    doc.add_paragraph("Data: {{ data_emissao }}")

    doc.add_heading("1. Objeto", 1)
    doc.add_paragraph("{{ objeto }}")

    doc.add_heading("2. Contexto e Justificativa", 1)
    doc.add_paragraph("{{ contexto }}")

    doc.add_heading("3. Escopo", 1)
    doc.add_paragraph("{%p for item in escopo %}")
    doc.add_paragraph("{{ item }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("4. Metodologia", 1)
    doc.add_paragraph("{{ metodologia }}")

    doc.add_heading("5. Etapas e Prazos", 1)
    doc.add_paragraph("{%p for e in etapas %}")
    doc.add_paragraph("{{ e.nome }} — {{ e.prazo }}", style="List Number")
    doc.add_paragraph("{{ e.descricao }}")
    doc.add_paragraph("{%p endfor %}")
    doc.add_paragraph("Prazo total: {{ prazo_total }}")

    doc.add_heading("6. Equipe Técnica", 1)
    doc.add_paragraph("{%p for m in equipe %}")
    doc.add_paragraph("{{ m.quantidade }}x {{ m.perfil }}: {{ m.atribuicoes }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("7. Atendimento aos Requisitos do TR", 1)
    doc.add_paragraph("{%p for r in requisitos_atendidos %}")
    doc.add_paragraph("{{ r }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("8. Condições de Pagamento", 1)
    doc.add_paragraph("{{ criterios_pagamento }}")

    doc.add_heading("9. Premissas", 1)
    doc.add_paragraph("{%p for p in premissas %}")
    doc.add_paragraph("{{ p }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("10. Riscos e Mitigações", 1)
    doc.add_paragraph("{%p for r in riscos %}")
    doc.add_paragraph("{{ r.risco }} — Mitigação: {{ r.mitigacao }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.save(str(TEMPLATES / "proposta.docx"))


def build_resumo():
    doc = _base_doc()
    doc.add_heading("RESUMO — {{ titulo }}", 0)
    doc.add_paragraph("Data: {{ data_emissao }}")
    doc.add_paragraph("Cliente: {{ cliente }}")
    doc.add_paragraph("Objeto: {{ objeto }}")
    doc.add_paragraph("Prazo: {{ prazo }}")
    doc.add_paragraph("Valor de referência: {{ valor_referencia }}")

    doc.add_heading("Escopo Principal", 1)
    doc.add_paragraph("{%p for item in escopo_resumido %}")
    doc.add_paragraph("{{ item }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("Destaques", 1)
    doc.add_paragraph("{%p for d in destaques %}")
    doc.add_paragraph("{{ d }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.add_heading("Próximos Passos", 1)
    doc.add_paragraph("{%p for p in proximos_passos %}")
    doc.add_paragraph("{{ p }}", style="List Bullet")
    doc.add_paragraph("{%p endfor %}")

    doc.save(str(TEMPLATES / "resumo.docx"))


if __name__ == "__main__":
    build_proposta()
    build_resumo()
    print(f"Templates gerados em {TEMPLATES}")
